"""Parser, phone, VCF, streaming, resume, and large-file tests."""

from __future__ import annotations

import csv
import threading
import time
from pathlib import Path

import pytest
from openpyxl import Workbook

from core.checkpoint import load_checkpoint
from core.control import RunControl
from exporters.vcf_validator import VcfStructureError, format_card, validate_vcf
from models.records import JobConfig
from parsers.csv_parser import CsvReader
from parsers.detect import inspect_source
from parsers.txt_parser import TxtReader
from parsers.xlsx_parser import XlsxReader
from processors.phone_normalizer import normalize_phone, prepare_name
from processors.streaming_engine import execute


def _config(path: Path, output: Path, **overrides: object) -> JobConfig:
    values: dict[str, object] = {
        "input_path": path,
        "output_dir": output,
        "file_format": path.suffix.lower().lstrip("."),
        "delimiter": "," if path.suffix.lower() == ".csv" else "|",
        "has_header": path.suffix.lower() != ".txt",
        "name_column": 0,
        "phone_column": 1,
        "contacts_per_file": 5000,
        "dedupe": True,
        "normalize_phone": True,
        "vn_to_e164": True,
        "keep_original_name": True,
        "split_directories": False,
        "encoding": "utf-8-sig",
        "chunk_size": 5000,
    }
    values.update(overrides)
    return JobConfig(**values)  # type: ignore[arg-type]


def _write_csv(path: Path, rows: list[tuple[str, str]], header: bool = True) -> None:
    with path.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.writer(handle)
        if header:
            writer.writerow(["name", "phone"])
        writer.writerows(rows)


def _cards(output: Path) -> list[tuple[str, int]]:
    found: list[tuple[str, int]] = []
    for path in sorted(output.rglob("contacts_*.vcf")):
        text = path.read_text(encoding="utf-8")
        found.append((str(path.relative_to(output)), text.count("BEGIN:VCARD")))
    return found


def _phones(output: Path) -> list[str]:
    phones: list[str] = []
    for path in sorted(output.rglob("contacts_*.vcf")):
        for line in path.read_text(encoding="utf-8").splitlines():
            if line.startswith("TEL;"):
                phones.append(line.split(":", 1)[1])
    return phones


def _names(output: Path) -> list[str]:
    names: list[str] = []
    for path in sorted(output.rglob("contacts_*.vcf")):
        logical = path.read_text(encoding="utf-8").replace("\r\n", "\n")
        folded: list[str] = []
        for line in logical.split("\n"):
            if folded and line.startswith(" "):
                folded[-1] += line[1:]
            else:
                folded.append(line)
        for line in folded:
            if line.startswith("FN:"):
                names.append(line[3:].replace("\\n", "\n").replace("\\;", ";").replace("\\,", ",").replace("\\\\", "\\"))
    return names


def test_csv_parser_reads_rows_without_assuming_only_english_headers(tmp_path: Path) -> None:
    path = tmp_path / "people.csv"
    _write_csv(path, [("Nguyen Van A", "0901234567"), ("Tran Thi B", "0912345678")])
    reader = CsvReader(path, ",", True, "utf-8-sig")
    assert reader.column_names() == ["name", "phone"]
    rows = list(reader.iter_stream(0, 1))
    assert [row.columns for row in rows] == [
        ("Nguyen Van A", "0901234567"),
        ("Tran Thi B", "0912345678"),
    ]
    assert rows[0].row_number == 2


def test_csv_keeps_newline_inside_quotes(tmp_path: Path) -> None:
    path = tmp_path / "quoted.csv"
    path.write_text('name,phone\n"Nguyen\nVan A",0901234567\n', encoding="utf-8")
    reader = CsvReader(path, ",", True, "utf-8-sig")
    rows = list(reader.iter_stream(0, 1))
    assert len(rows) == 1
    assert rows[0].columns[0] == "Nguyen\nVan A"


def test_xlsx_parser(tmp_path: Path) -> None:
    path = tmp_path / "people.xlsx"
    book = Workbook()
    sheet = book.active
    sheet.append(["name", "phone"])
    sheet.append(["Nguyen Van A", "0901234567"])
    book.save(path)
    reader = XlsxReader(path, True)
    assert reader.column_names() == ["name", "phone"]
    rows = list(reader.iter_stream(0, 1))
    assert rows[0].columns == ("Nguyen Van A", "0901234567")


def test_xlsx_numeric_phone_is_not_guessed(tmp_path: Path) -> None:
    path = tmp_path / "numbers.xlsx"
    book = Workbook()
    sheet = book.active
    sheet.append(["name", "phone"])
    sheet.append(["Nguyen Van A", 901234567])
    book.save(path)
    result = execute(_config(path, tmp_path / "out"))
    assert result.report is not None
    assert result.report.invalid == 1
    assert result.report.exported == 0
    errors = (tmp_path / "out" / "errors.csv").read_text(encoding="utf-8-sig")
    assert "không đủ cơ sở" in errors


def test_txt_parser_and_custom_delimiter(tmp_path: Path) -> None:
    path = tmp_path / "people.txt"
    path.write_text("Nguyen Van A|0901234567\nTran Thi B;0912345678\n", encoding="utf-8")
    pipe = TxtReader(path, "|", False, "utf-8-sig")
    rows = list(pipe.iter_stream(0, 1))
    assert rows[0].columns[0] == "Nguyen Van A"
    assert rows[0].columns[1] == "0901234567"
    semi = TxtReader(path, ";", False, "utf-8-sig")
    assert list(semi.iter_stream(0, 1))[1].columns[0] == "Tran Thi B"


def test_phone_vietnam_international_and_spaces() -> None:
    domestic = normalize_phone("0901234567", normalize=True, vn_to_e164=True)
    plus = normalize_phone("+84901234567", normalize=True, vn_to_e164=True)
    spaced = normalize_phone("090 123 4567", normalize=True, vn_to_e164=True)
    dotted = normalize_phone("090.123.4567", normalize=True, vn_to_e164=True)
    paren = normalize_phone("(090) 123-4567", normalize=True, vn_to_e164=True)
    assert domestic.phone == "+84901234567"
    assert plus.phone == "+84901234567"
    assert spaced.phone == "+84901234567"
    assert dotted.phone == "+84901234567"
    assert paren.phone == "+84901234567"


def test_phone_without_enough_evidence_is_rejected() -> None:
    short = normalize_phone("12345", normalize=True, vn_to_e164=True)
    letters = normalize_phone("090abc4567", normalize=True, vn_to_e164=True)
    bare = normalize_phone("901234567", normalize=True, vn_to_e164=True)
    assert not short.ok
    assert not letters.ok
    assert not bare.ok
    assert "không đủ cơ sở" in bare.reason


def test_phone_keeps_original_when_normalize_is_off() -> None:
    result = normalize_phone("090 123 4567", normalize=False, vn_to_e164=True)
    assert result.ok
    assert result.phone == "090 123 4567"


def test_vietnamese_and_unicode_names_round_trip(tmp_path: Path) -> None:
    path = tmp_path / "names.csv"
    rows = [
        ("Nguyễn Văn Á", "0901234567"),
        ("山田太郎", "0901234568"),
        ("Ten\nEND:VCARD\nBEGIN:VCARD", "0901234569"),
        ("Nguyen; Van, A", "0901234570"),
    ]
    _write_csv(path, rows)
    output = tmp_path / "out"
    result = execute(_config(path, output, dedupe=False))
    assert result.report is not None
    assert result.report.exported == 4
    names = _names(output)
    assert names[0] == "Nguyễn Văn Á"
    assert names[1] == "山田太郎"
    assert "END:VCARD" in names[2]
    assert names[3] == "Nguyen; Van, A"
    for vcf in output.glob("contacts_*.vcf"):
        validate_vcf(vcf, 4)
        lines = [
            line.rstrip("\r")
            for line in vcf.read_bytes().decode("utf-8").split("\n")
        ]
        assert lines.count("BEGIN:VCARD") == lines.count("END:VCARD") == 4


def test_keep_original_name_preserves_inner_spaces(tmp_path: Path) -> None:
    path = tmp_path / "spaces.csv"
    _write_csv(path, [("  Nguyen   Van A  ", "0901234567")])
    kept = tmp_path / "kept"
    execute(_config(path, kept, keep_original_name=True))
    assert _names(kept) == ["  Nguyen   Van A  "]
    collapsed = tmp_path / "collapsed"
    execute(_config(path, collapsed, keep_original_name=False))
    assert _names(collapsed) == ["Nguyen Van A"]


def test_duplicate_phones_are_removed_across_the_whole_run(tmp_path: Path) -> None:
    path = tmp_path / "dup.csv"
    _write_csv(
        path,
        [
            ("Mot", "0901234567"),
            ("Hai", "0912345678"),
            ("Ba", "090 123 4567"),
            ("Bon", "+84901234567"),
        ],
    )
    output = tmp_path / "out"
    result = execute(_config(path, output, contacts_per_file=2, chunk_size=2))
    assert result.report is not None
    assert result.report.duplicate == 2
    assert result.report.exported == 2
    assert _phones(output) == ["+84901234567", "+84912345678"]


def test_invalid_row_does_not_stop_the_file(tmp_path: Path) -> None:
    path = tmp_path / "mixed.csv"
    _write_csv(
        path,
        [
            ("Mot", "0901234567"),
            ("Hong", "abc"),
            ("", "0901234568"),
            ("Ba", "0901234569"),
        ],
    )
    output = tmp_path / "out"
    result = execute(_config(path, output))
    assert result.report is not None
    assert result.report.total_rows == 4
    assert result.report.invalid == 2
    assert result.report.exported == 2
    errors = list(csv.DictReader((output / "errors.csv").open(encoding="utf-8-sig")))
    assert [row["error_reason"] for row in errors] == [
        "Số chứa ký tự không hợp lệ",
        "Thiếu tên",
    ]
    assert errors[0]["row_number"] == "3"


@pytest.mark.parametrize(
    ("count", "expected"),
    [
        (4999, [4999]),
        (5000, [5000]),
        (5001, [5000, 1]),
    ],
)
def test_file_split_boundaries(tmp_path: Path, count: int, expected: list[int]) -> None:
    path = tmp_path / "people.csv"
    _write_csv(path, [(f"N{index}", f"09{index:08d}") for index in range(count)])
    output = tmp_path / "out"
    result = execute(_config(path, output, contacts_per_file=5000, chunk_size=1000))
    assert result.report is not None
    assert result.report.exported == count
    counts = [item[1] for item in _cards(output)]
    assert counts == expected
    for vcf, size in zip(sorted(output.glob("contacts_*.vcf")), expected, strict=True):
        validate_vcf(vcf, size)
        assert vcf.name.startswith("contacts_")


def test_empty_file_creates_no_contact(tmp_path: Path) -> None:
    path = tmp_path / "empty.csv"
    path.write_text("name,phone\n", encoding="utf-8")
    output = tmp_path / "out"
    result = execute(_config(path, output))
    assert result.report is not None
    assert result.report.total_rows == 0
    assert result.report.vcf_files == 0
    assert _cards(output) == []
    report = (output / "report.txt").read_text(encoding="utf-8")
    assert "TOTAL ROWS: 0" in report
    assert "VCF FILES: 0" in report


def test_split_directories(tmp_path: Path) -> None:
    path = tmp_path / "people.csv"
    _write_csv(path, [(f"N{index}", f"09{index:08d}") for index in range(5)])
    output = tmp_path / "out"
    execute(
        _config(
            path,
            output,
            contacts_per_file=2,
            split_directories=True,
            files_per_batch=2,
            chunk_size=2,
        )
    )
    assert (output / "batch_0001" / "contacts_00001.vcf").is_file()
    assert (output / "batch_0001" / "contacts_00002.vcf").is_file()
    assert (output / "batch_0002" / "contacts_00003.vcf").is_file()
    assert sum(item[1] for item in _cards(output)) == 5


def test_txt_conversion(tmp_path: Path) -> None:
    path = tmp_path / "people.txt"
    path.write_text("Nguyen Van A|0901234567\n", encoding="utf-8")
    output = tmp_path / "out"
    result = execute(_config(path, output, file_format="txt", delimiter="|", has_header=False))
    assert result.report is not None
    assert result.report.exported == 1
    assert _phones(output) == ["+84901234567"]
    raw = next(output.glob("contacts_*.vcf")).read_bytes()
    assert raw.startswith(b"BEGIN:VCARD\r\n")
    assert b"END:VCARD\r\n" in raw


def test_resume_after_cancel_does_not_restart(tmp_path: Path) -> None:
    path = tmp_path / "people.csv"
    _write_csv(path, [(f"N{index:04d}", f"09{index:08d}") for index in range(300)])
    output = tmp_path / "out"

    class StopAfter(RunControl):
        def observe(self, progress: object) -> None:
            processed = getattr(progress, "processed", 0)
            if processed >= 120:
                self.cancel()

    stopped = execute(
        _config(path, output, contacts_per_file=100, chunk_size=50),
        control=StopAfter(),
    )
    assert stopped.cancelled
    checkpoint = load_checkpoint(output)
    assert checkpoint is not None
    assert checkpoint.processed == 100
    assert checkpoint.exported == 100

    resumed = execute(
        _config(path, output, contacts_per_file=100, chunk_size=50),
        resume=True,
    )
    assert resumed.report is not None
    assert resumed.report.exported == 300
    assert resumed.report.total_rows == 300
    assert [item[1] for item in _cards(output)] == [100, 100, 100]
    assert len(set(_phones(output))) == 300


def test_resume_rejects_a_changed_source_file(tmp_path: Path) -> None:
    path = tmp_path / "people.csv"
    _write_csv(path, [(f"N{index}", f"09{index:08d}") for index in range(30)])
    output = tmp_path / "out"

    class StopAfter(RunControl):
        def observe(self, progress: object) -> None:
            if getattr(progress, "processed", 0) >= 10:
                self.cancel()

    execute(_config(path, output, contacts_per_file=100, chunk_size=10), control=StopAfter())
    path.write_text(path.read_text(encoding="utf-8") + "Them,0900000000\n", encoding="utf-8")
    with pytest.raises(ValueError, match="File nguồn đã thay đổi"):
        execute(_config(path, output, contacts_per_file=100, chunk_size=10), resume=True)


def test_pause_blocks_until_resume(tmp_path: Path) -> None:
    path = tmp_path / "people.csv"
    _write_csv(path, [("Mot", "0901234567"), ("Hai", "0901234568")])
    output = tmp_path / "out"
    control = RunControl()
    control.pause()

    def release() -> None:
        time.sleep(0.2)
        control.resume()

    threading.Thread(target=release, daemon=True).start()
    started = time.monotonic()
    result = execute(_config(path, output), control=control)
    assert time.monotonic() - started >= 0.15
    assert result.report is not None
    assert result.report.exported == 2


def test_large_streaming_run_keeps_exact_batches(tmp_path: Path) -> None:
    count = 20_001
    path = tmp_path / "large.csv"
    with path.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.writer(handle)
        writer.writerow(["name", "phone"])
        for index in range(count):
            writer.writerow((f"N{index}", f"09{index:08d}"))
    output = tmp_path / "out"
    result = execute(_config(path, output, contacts_per_file=5000, chunk_size=2000))
    assert result.report is not None
    assert result.report.exported == count
    assert result.report.invalid == 0
    counts = [item[1] for item in _cards(output)]
    assert counts == [5000, 5000, 5000, 5000, 1]
    validate_vcf(output / "contacts_00001.vcf", 5000)
    validate_vcf(output / "contacts_00005.vcf", 1)


def test_reader_is_a_generator() -> None:
    import inspect

    assert inspect.isgeneratorfunction(CsvReader.iter_stream)
    assert inspect.isgeneratorfunction(TxtReader.iter_stream)
    assert inspect.isgeneratorfunction(XlsxReader.iter_stream)


def test_broken_vcf_is_rejected(tmp_path: Path) -> None:
    path = tmp_path / "contacts_00001.vcf"
    path.write_text("BEGIN:VCARD\r\nVERSION:3.0\r\nFN:A\r\nTEL;TYPE=CELL:+84901234567\r\n", encoding="utf-8")
    with pytest.raises(VcfStructureError):
        validate_vcf(path, 1)


def test_prepare_name_rejects_blank() -> None:
    name, reason = prepare_name("   ", keep_original=True)
    assert name is None
    assert reason == "Thiếu tên"


def test_inspect_csv_columns(tmp_path: Path) -> None:
    path = tmp_path / "people.csv"
    _write_csv(path, [("An", "0901234567")])
    info = inspect_source(path, delimiter=",", has_header=True)
    assert info.columns == ["name", "phone"]
    assert info.estimated_rows == 1
    assert info.samples[0][0] == "An"


def test_source_file_is_not_deleted(tmp_path: Path) -> None:
    path = tmp_path / "people.csv"
    _write_csv(path, [("An", "0901234567")])
    before = path.read_bytes()
    execute(_config(path, tmp_path / "out"))
    assert path.read_bytes() == before


def test_window_loads_columns_without_blocking(tmp_path: Path) -> None:
    from PySide6.QtWidgets import QApplication

    from ui.main_window import MainWindow

    path = tmp_path / "people.csv"
    _write_csv(path, [("Nguyen Van A", "0901234567")])
    app = QApplication.instance() or QApplication([])
    window = MainWindow()
    window.file_edit.setText(str(path))
    window.delimiter_combo.setCurrentIndex(0)
    window.header_check.setChecked(True)
    window._reinspect()
    deadline = time.monotonic() + 5
    while window.inspection is None and time.monotonic() < deadline:
        app.processEvents()
        time.sleep(0.02)
    assert window.inspection is not None
    assert window.name_combo.itemText(0) == "name"
    assert window.phone_combo.itemText(1) == "phone"
    assert window.info_labels["format"].text() == "CSV"
    window.delimiter_combo.setCurrentIndex(1)
    window.delimiter_combo.setCurrentIndex(0)
    window.close()
    app.processEvents()


def test_xlsx_blank_rows_are_skipped(tmp_path: Path) -> None:
    path = tmp_path / "people.xlsx"
    book = Workbook()
    sheet = book.active
    sheet.append(["name", "phone"])
    sheet.append(["Mot", "0901234567"])
    sheet.append([None, None])
    sheet.append(["Hai", "0901234568"])
    book.save(path)
    rows = list(XlsxReader(path, True).iter_stream(0, 1))
    assert [row.columns[0] for row in rows] == ["Mot", "Hai"]


def test_xlsx_resume_after_cancel(tmp_path: Path) -> None:
    path = tmp_path / "people.xlsx"
    book = Workbook()
    sheet = book.active
    sheet.append(["name", "phone"])
    for index in range(25):
        sheet.append([f"N{index:02d}", f"09{index:08d}"])
    book.save(path)
    output = tmp_path / "out"

    class StopAfter(RunControl):
        def observe(self, progress: object) -> None:
            if getattr(progress, "processed", 0) >= 12:
                self.cancel()

    stopped = execute(
        _config(path, output, contacts_per_file=10, chunk_size=5),
        control=StopAfter(),
    )
    assert stopped.cancelled
    checkpoint = load_checkpoint(output)
    assert checkpoint is not None
    assert checkpoint.processed == 10

    resumed = execute(
        _config(path, output, delimiter=",", contacts_per_file=10, chunk_size=5),
        resume=True,
    )
    assert resumed.report is not None
    assert resumed.report.exported == 25
    assert load_checkpoint(output) is None


def test_checkpoint_in_sqlite_wins_over_a_stale_json_file(tmp_path: Path) -> None:
    path = tmp_path / "people.csv"
    _write_csv(path, [(f"N{index:04d}", f"09{index:08d}") for index in range(40)])
    output = tmp_path / "out"

    class StopAfter(RunControl):
        def observe(self, progress: object) -> None:
            if getattr(progress, "processed", 0) >= 20:
                self.cancel()

    execute(_config(path, output, contacts_per_file=100, chunk_size=10), control=StopAfter())
    saved = load_checkpoint(output)
    assert saved is not None
    assert saved.processed == 20
    json_path = output / ".contact_to_vcf" / "checkpoint.json"
    payload = json_path.read_text(encoding="utf-8")
    json_path.write_text(payload.replace('"processed": 20', '"processed": 1'), encoding="utf-8")
    reloaded = load_checkpoint(output)
    assert reloaded is not None
    assert reloaded.processed == 20


def test_finished_job_drops_resume_state(tmp_path: Path) -> None:
    path = tmp_path / "people.csv"
    _write_csv(path, [("An", "0901234567")])
    output = tmp_path / "out"
    result = execute(_config(path, output))
    assert result.report is not None
    assert (output / "contacts_00001.vcf").is_file()
    assert (output / "report.txt").is_file()
    assert load_checkpoint(output) is None
    assert not (output / ".contact_to_vcf").exists()


def test_window_restores_saved_options(tmp_path: Path) -> None:
    from PySide6.QtWidgets import QApplication

    from ui.main_window import MainWindow

    path = tmp_path / "people.csv"
    _write_csv(path, [(f"N{index}", f"09{index:08d}") for index in range(30)])
    output = tmp_path / "out"

    class StopAfter(RunControl):
        def observe(self, progress: object) -> None:
            if getattr(progress, "processed", 0) >= 10:
                self.cancel()

    execute(
        _config(path, output, contacts_per_file=80, dedupe=False, chunk_size=10),
        control=StopAfter(),
    )
    app = QApplication.instance() or QApplication([])
    window = MainWindow()
    window.file_edit.setText(str(path))
    window.delimiter_combo.setCurrentIndex(0)
    window.header_check.setChecked(True)
    window._reinspect()
    deadline = time.monotonic() + 5
    while window.inspection is None and time.monotonic() < deadline:
        app.processEvents()
        time.sleep(0.02)
    window._use_output(output)
    app.processEvents()
    assert window.per_file.value() == 80
    assert window.dedupe_check.isChecked() is False
    assert window.resume_btn.isEnabled()
    assert window.name_combo.currentIndex() == 0
    assert window.phone_combo.currentIndex() == 1
    window.close()
    app.processEvents()


def test_several_phones_in_one_cell_become_separate_contacts(tmp_path: Path) -> None:
    path = tmp_path / "multi.csv"
    _write_csv(
        path,
        [
            ("An", "0901234567 / 0912345678"),
            ("Binh", "0901234569, 0987654321"),
            ("Chi", "090 123 4570"),
            ("Dung", "0901111111 / khong-phai-so"),
        ],
    )
    output = tmp_path / "out"
    result = execute(_config(path, output, contacts_per_file=2))
    assert result.report is not None
    assert result.report.total_rows == 4
    assert result.report.exported == 6
    assert result.report.invalid == 0
    assert _phones(output) == [
        "+84901234567",
        "+84912345678",
        "+84901234569",
        "+84987654321",
        "+84901234570",
        "+84901111111",
    ]
    guide = (output / "thu_tu_nhap.txt").read_text(encoding="utf-8")
    assert "contacts_00001.vcf — 2 liên hệ — An → An" in guide
    assert "contacts_00003.vcf — 2 liên hệ — Chi → Dung" in guide
    assert "Thêm vào Danh bạ" in guide


def test_semicolon_csv_is_detected_without_a_fixed_delimiter(tmp_path: Path) -> None:
    path = tmp_path / "people.csv"
    path.write_text("ten;so\nAn;0901234567\n", encoding="utf-8")
    info = inspect_source(path)
    assert info.delimiter == ";"
    assert info.columns == ["ten", "so"]
    assert info.samples[0] == ("An", "0901234567")


def test_phone_column_is_suggested_from_sample_values() -> None:
    from parsers.column_suggest import suggest_columns

    name_index, phone_index = suggest_columns(
        ["so", "ten"],
        [("0901234567", "An"), ("0912345678", "Binh")],
    )
    assert phone_index == 0
    assert name_index == 1


def test_window_shows_preview_and_suggested_phone_column(tmp_path: Path) -> None:
    from PySide6.QtWidgets import QApplication

    from ui.main_window import MainWindow

    path = tmp_path / "people.csv"
    path.write_text("so;ten\n0901234567;An\n", encoding="utf-8")
    app = QApplication.instance() or QApplication([])
    window = MainWindow()
    window.file_edit.setText(str(path))
    window.header_check.setChecked(True)
    window._auto_delimiter = True
    window._reinspect()
    deadline = time.monotonic() + 5
    while window.inspection is None and time.monotonic() < deadline:
        app.processEvents()
        time.sleep(0.02)
    assert window.inspection is not None
    assert window.inspection.delimiter == ";"
    assert window.phone_combo.currentIndex() == 0
    assert window.name_combo.currentIndex() == 1
    assert window.preview.rowCount() == 1
    assert window.preview.item(0, 0).text() == "0901234567"
    assert "Cột số điện thoại: so" in window.mapping_label.text()
    window.close()
    app.processEvents()


def test_pool_keeps_a_number_in_its_first_book(tmp_path: Path) -> None:
    from processors.phone_normalizer import canonical_phones
    from processors.pool import book_of, export_book, import_file, list_books, pool_total

    assert canonical_phones("090 123 4567") == ["+84901234567"]
    assert canonical_phones("+84 901 234 567") == ["+84901234567"]
    assert canonical_phones("12") == ["12"]
    assert canonical_phones("abc") == []
    assert canonical_phones("0901234567 / 0912345678") == ["+84901234567", "+84912345678"]

    folder = tmp_path / "kho"
    folder.mkdir()
    source = tmp_path / "a.csv"
    _write_csv(
        source,
        [
            ("Bo qua", "0901234567"),
            ("Ngan", "12"),
            ("Trong", "abc"),
            ("Hai so", "0901234568 / 0987654321"),
        ],
    )
    first = import_file(
        folder,
        source,
        file_format="csv",
        delimiter=",",
        has_header=True,
        encoding="utf-8-sig",
        phone_column=1,
        contacts_per_file=2,
    )
    assert first.added == 4
    assert first.rejected == 1
    assert pool_total(folder) == 4
    assert [book.contact_count for book in list_books(folder)] == [2, 2]
    home = book_of(folder, "+84901234567")
    assert home == 1

    again = import_file(
        folder,
        source,
        file_format="csv",
        delimiter=",",
        has_header=True,
        encoding="utf-8-sig",
        phone_column=1,
        contacts_per_file=2,
    )
    assert again.added == 0
    assert again.duplicate == 4
    assert book_of(folder, "+84901234567") == home
    assert [book.contact_count for book in list_books(folder)] == [2, 2]

    destination = tmp_path / "danhba_00001.vcf"
    assert export_book(folder, 1, destination) == 2
    text = destination.read_text(encoding="utf-8")
    assert "FN:0901234567" in text
    assert "TEL;TYPE=CELL:+84901234567" in text
    assert text.count("BEGIN:VCARD") == 2
    downloaded = list_books(folder, "downloaded")
    assert downloaded[0].id == 1
    assert downloaded[0].downloaded_at

    extra = tmp_path / "b.csv"
    _write_csv(extra, [("Them", "0901234569")])
    import_file(
        folder,
        extra,
        file_format="csv",
        delimiter=",",
        has_header=True,
        encoding="utf-8-sig",
        phone_column=1,
        contacts_per_file=2,
    )
    assert book_of(folder, "+84901234567") == 1
    assert book_of(folder, "+84901234569") == 3
    assert [book.id for book in list_books(folder, "pending")] == [2, 3]


def test_card_bytes_match_the_working_phone_sample() -> None:
    sample_path = Path(__file__).resolve().parents[1] / "src" / "web" / "test_10_contacts.vcf"
    sample = sample_path.read_bytes()
    built = b"".join(
        format_card(f"Test {index:02d}", f"+8490000{index:04d}").encode("utf-8")
        for index in range(1, 11)
    )
    assert built == sample
    thirty = (sample_path.parent / "danhba_test_30.vcf").read_bytes()
    assert thirty.startswith(built)
    assert thirty.count(b"BEGIN:VCARD") == 30
    assert b"TEL;TYPE=CELL:+84900000030\r\n" in thirty
    assert b"0900000001" not in thirty


def test_window_constructs() -> None:
    from PySide6.QtWidgets import QApplication

    from ui.main_window import MainWindow

    app = QApplication.instance() or QApplication([])
    window = MainWindow()
    assert window.windowTitle() == "Chuyển danh bạ sang VCF"
    assert window.per_file.value() == 5000
    window.close()
    assert app is not None
