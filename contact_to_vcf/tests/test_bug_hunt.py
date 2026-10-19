"""Edge cases that used to drop phones or close a book too early."""

from __future__ import annotations

import json
from http.cookiejar import CookieJar
from pathlib import Path
from urllib.error import HTTPError
from urllib.request import HTTPCookieProcessor, Request, build_opener

from openpyxl import Workbook

from parsers.column_suggest import phone_column_for_import
from parsers.detect import inspect_source
from parsers.xlsx_parser import XlsxReader
from processors.phone_normalizer import canonical_phones, to_vietnam_10
from processors.pool import (
    book_of,
    import_file,
    list_books,
    pool_total,
    read_book_vcf,
)
from web.server import serve


def test_excel_text_and_country_code_become_ten_digits() -> None:
    samples = {
        "0901234567.0": "0901234567",
        "84901234567.0": "0901234567",
        "8.4901234567E+10": "0901234567",
        "8401621234567": "0321234567",
        "+84 0162 123 4567": "0321234567",
        "０９０１２３４５６７": "0901234567",
        "'0901234567": "0901234567",
        "090.123.4567": "0901234567",
        "84 901 234 567": "0901234567",
    }
    for raw, expected in samples.items():
        assert to_vietnam_10(raw) == expected, raw
        assert canonical_phones(raw) == [expected], raw


def test_excel_uses_column_b_and_keeps_the_first_data_row(tmp_path: Path) -> None:
    path = tmp_path / "so.xlsx"
    book = Workbook()
    sheet = book.active
    assert sheet is not None
    sheet.append(["123456789012", 84901234567])
    sheet.append(["123456789013", "090.123.4568"])
    second = book.create_sheet("them")
    second.append(["ma", "sdt"])
    second.append(["x", "0912345678"])
    book.save(path)

    rows = list(XlsxReader(path, True).iter_stream(0, 1))
    assert rows[0].columns[1] == "84901234567"
    info = inspect_source(path)
    assert info.has_header is False
    assert phone_column_for_import(info.file_format, info.columns, info.samples) == 1

    folder = tmp_path / "kho"
    folder.mkdir()
    stats = import_file(
        folder,
        path,
        file_format="xlsx",
        delimiter="",
        has_header=info.has_header,
        encoding="utf-8-sig",
        phone_column=1,
        contacts_per_file=1000,
    )
    assert stats.added == 3
    assert stats.filled == ()
    assert pool_total(folder) == 3
    assert book_of(folder, "0901234567") == 1
    assert book_of(folder, "0901234568") == 1
    assert book_of(folder, "0912345678") == 1
    assert book_of(folder, "123456789012") is None
    books = list_books(folder)
    assert books[0].contact_count == 3
    assert books[0].downloaded_at is None
    read_book_vcf(folder, 1)
    assert list_books(folder)[0].downloaded_at is None


def test_open_book_keeps_accepting_numbers_until_one_thousand(tmp_path: Path) -> None:
    folder = tmp_path / "kho"
    folder.mkdir()
    first = tmp_path / "a.txt"
    first.write_text("\n".join(f"090{index:07d}" for index in range(600)), encoding="utf-8")
    second = tmp_path / "b.txt"
    second.write_text("\n".join(f"091{index:07d}" for index in range(500)), encoding="utf-8")
    added = import_file(
        folder,
        first,
        file_format="txt",
        delimiter=",",
        has_header=False,
        encoding="utf-8",
        phone_column=0,
    )
    assert added.added == 600
    assert added.filled == ()
    assert [book.contact_count for book in list_books(folder)] == [600]
    more = import_file(
        folder,
        second,
        file_format="txt",
        delimiter=",",
        has_header=False,
        encoding="utf-8",
        phone_column=0,
    )
    assert more.added == 500
    assert more.filled == (1,)
    counts = [book.contact_count for book in list_books(folder)]
    assert counts == [1000, 100]
    assert list_books(folder, "downloaded") == []
    again = import_file(
        folder,
        first,
        file_format="txt",
        delimiter=",",
        has_header=False,
        encoding="utf-8",
        phone_column=0,
    )
    assert again.added == 0
    assert again.duplicate == 600
    assert [book.contact_count for book in list_books(folder)] == [1000, 100]
    assert book_of(folder, "0900000000") == 1
    assert book_of(folder, "0910000000") == 1
    assert book_of(folder, "0910000400") == 2


def test_browser_import_keeps_column_b_and_does_not_freeze_a_short_book(tmp_path: Path) -> None:
    path = tmp_path / "ds.xlsx"
    book = Workbook()
    sheet = book.active
    assert sheet is not None
    sheet.append(["ma khach", "dien thoai"])
    sheet.append(["KH-1", "84912345678"])
    sheet.append(["KH-2", "0901234567.0"])
    book.save(path)

    server = serve(tmp_path / "kho", host="127.0.0.1", port=0)
    import threading

    threading.Thread(target=server.serve_forever, daemon=True).start()
    port = server.server_address[1]
    opener = build_opener(HTTPCookieProcessor(CookieJar()))
    try:
        opener.open(
            Request(
                f"http://127.0.0.1:{port}/api/login",
                data=json.dumps({"user": "danhba", "password": "danhba123"}).encode(),
                headers={"Content-Type": "application/json"},
            )
        )
        body, content_type = _form(("danh-sach.xlsx", path.read_bytes()))
        imported = json.load(
            opener.open(
                Request(
                    f"http://127.0.0.1:{port}/api/import-now",
                    data=body,
                    headers={"Content-Type": content_type},
                )
            )
        )
        assert imported["added"] == 2
        assert imported["phone_column"] == 1
        assert imported["filled"] == []
        books = json.load(opener.open(f"http://127.0.0.1:{port}/api/books"))
        assert books["books"][0]["count"] == 2
        assert books["books"][0]["downloaded_at"] is None
        card = opener.open(
            f"http://127.0.0.1:{port}/api/books/{books['books'][0]['id']}/download"
        ).read().decode()
        assert "FN:0912345678" in card
        assert "TEL;TYPE=CELL:0901234567" in card
        assert "123456789" not in card
        frozen = json.load(opener.open(f"http://127.0.0.1:{port}/api/books?status=downloaded"))
        assert frozen["books"][0]["downloaded_at"]
    finally:
        server.shutdown()


def test_broken_excel_is_rejected_without_crashing(tmp_path: Path) -> None:
    server = serve(tmp_path / "kho", host="127.0.0.1", port=0)
    import threading

    threading.Thread(target=server.serve_forever, daemon=True).start()
    port = server.server_address[1]
    opener = build_opener(HTTPCookieProcessor(CookieJar()))
    try:
        opener.open(
            Request(
                f"http://127.0.0.1:{port}/api/login",
                data=json.dumps({"user": "danhba", "password": "danhba123"}).encode(),
                headers={"Content-Type": "application/json"},
            )
        )
        body, content_type = _form(("hong.xlsx", b"day khong phai excel"))
        try:
            opener.open(
                Request(
                    f"http://127.0.0.1:{port}/api/import-now",
                    data=body,
                    headers={"Content-Type": content_type},
                )
            )
        except HTTPError as exc:
            payload = json.load(exc)
            assert exc.code == 400
            assert payload["error"]
        else:
            raise AssertionError("file hỏng phải bị từ chối")
        books = json.load(opener.open(f"http://127.0.0.1:{port}/api/books"))
        assert books["total"] == 0
    finally:
        server.shutdown()


def _form(file_item: tuple[str, bytes]) -> tuple[bytes, str]:
    filename, payload = file_item
    boundary = "----bug-hunt"
    body = (
        f"--{boundary}\r\n"
        f'Content-Disposition: form-data; name="file"; filename="{filename}"\r\n'
        "Content-Type: application/octet-stream\r\n\r\n"
    ).encode() + payload + f"\r\n--{boundary}--\r\n".encode()
    return body, f"multipart/form-data; boundary={boundary}"
