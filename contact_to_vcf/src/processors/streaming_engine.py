"""Chunked conversion pipeline. Rows are never accumulated into one list."""

from __future__ import annotations

import json
import time
from pathlib import Path

from core.checkpoint import (
    CHECKPOINT_VERSION,
    clear_resume_state,
    config_matches,
    dedupe_path,
    load_checkpoint,
    save_checkpoint,
)
from core.control import RunControl
from core.output_reset import reset_output
from exporters.error_log import ErrorLog
from exporters.import_guide import write_import_guide
from exporters.vcf_exporter import VcfExporter
from exporters.vcf_validator import VcfStructureError, count_markers
from models.records import Checkpoint, JobConfig, JobResult, Progress, RawRecord, Report
from parsers.detect import open_reader
from processors.deduplicator import Deduplicator
from processors.row_processor import process_row
from utils.fingerprint import estimate_text_rows, file_fingerprint


def execute(
    config: JobConfig,
    *,
    mode: str = "convert",
    resume: bool = False,
    control: RunControl | None = None,
    on_progress=None,
    on_message=None,
) -> JobResult:
    if mode not in {"convert", "validate"}:
        raise ValueError("Chế độ xử lý không hợp lệ")
    config.validate()
    control = control or RunControl()
    converting = mode == "convert"
    if converting and not resume:
        reset_output(config.output_dir)
    config.output_dir.mkdir(parents=True, exist_ok=True)

    checkpoint = load_checkpoint(config.output_dir) if converting and resume else None
    if converting and resume:
        if checkpoint is None:
            raise ValueError("Không có tiến trình đã lưu để tiếp tục")
        mismatch = config_matches(checkpoint, config)
        if mismatch:
            raise ValueError(mismatch)

    size, fingerprint = file_fingerprint(config.input_path)
    total = checkpoint.total_rows if checkpoint and checkpoint.total_rows else _estimate(config)
    progress = Progress(total=total)
    if checkpoint is not None:
        progress.processed = checkpoint.processed
        progress.valid = checkpoint.valid
        progress.invalid = checkpoint.invalid
        progress.duplicate = checkpoint.duplicate
        progress.exported = checkpoint.exported
        progress.files_created = checkpoint.completed_files + (
            1 if checkpoint.current_file_contacts else 0
        )
        progress.elapsed_sec = checkpoint.elapsed_sec

    logger = _Logger(config.output_dir / "conversion.log", on_message if converting else on_message)
    deduper: Deduplicator | None = None
    errors: ErrorLog | None = None
    exporter: VcfExporter | None = None
    reader = open_reader(config)
    started = time.monotonic()
    base_elapsed = progress.elapsed_sec
    pending: list[tuple[int, str, str, str]] = []
    rows_since_commit = 0
    next_offset = 0 if checkpoint is None else checkpoint.next_offset
    next_row = 1 if checkpoint is None else checkpoint.next_row
    last_emit = 0.0
    cancelled = False
    finished_cleanly = False

    try:
        if converting:
            exporter = VcfExporter(
                config.output_dir,
                config.contacts_per_file,
                split_directories=config.split_directories,
                files_per_batch=config.files_per_batch,
            )
            if checkpoint is None:
                errors = ErrorLog(config.output_dir / "errors.csv", resume_bytes=None)
                if config.dedupe:
                    deduper = Deduplicator(dedupe_path(config.output_dir))
                    deduper.begin()
            else:
                if config.dedupe and not dedupe_path(config.output_dir).is_file():
                    raise ValueError(
                        "Thiếu dữ liệu lọc trùng của tiến trình đã lưu. Hãy bắt đầu lại."
                    )
                exporter.prepare_resume(
                    checkpoint.completed_files,
                    checkpoint.current_file_contacts,
                )
                errors = ErrorLog(
                    config.output_dir / "errors.csv",
                    resume_bytes=checkpoint.errors_bytes,
                )
                if config.dedupe:
                    deduper = Deduplicator(dedupe_path(config.output_dir))
                    deduper.begin()
            logger.write(
                "Tiếp tục chuyển đổi từ dòng đã lưu."
                if checkpoint
                else f"Bắt đầu chuyển đổi {config.input_path.name}."
            )
        elif config.dedupe:
            temporary = config.output_dir / ".contact_to_vcf" / "preview-dedupe.sqlite"
            if temporary.exists():
                temporary.unlink()
            deduper = Deduplicator(temporary)
            deduper.begin()
            logger.write("Bắt đầu kiểm tra dữ liệu.")

        def elapsed() -> float:
            return base_elapsed + (time.monotonic() - started)

        def emit(force: bool = False) -> None:
            nonlocal last_emit
            now = time.monotonic()
            if not force and now - last_emit < 0.5:
                return
            last_emit = now
            progress.elapsed_sec = elapsed()
            if progress.total is not None and progress.processed > progress.total:
                progress.total = progress.processed
            if on_progress is not None:
                on_progress(progress)

        def durable() -> None:
            if not converting or exporter is None or errors is None:
                if deduper is not None:
                    deduper.commit()
                    deduper.begin()
                return
            exporter.flush()
            error_bytes = errors.flush() if not pending else _flush_errors(errors, pending)
            pending.clear()
            progress.files_created = exporter.files_created
            progress.elapsed_sec = elapsed()
            checkpoint_now = _make_checkpoint(
                config,
                size,
                fingerprint,
                progress,
                exporter,
                error_bytes,
                next_offset,
                next_row,
            )
            if deduper is not None:
                deduper.save_state(
                    json.dumps(checkpoint_now.to_dict(), ensure_ascii=False)
                )
                deduper.commit()
                deduper.begin()
            save_checkpoint(checkpoint_now)

        emit(force=True)
        stream = reader.iter_stream(0 if checkpoint is None else checkpoint.next_offset, next_row)
        for record in stream:
            if not control.keep_going():
                cancelled = True
                break
            contacts, reason = process_row(record, config)
            file_closed = False
            if not contacts:
                progress.invalid += 1
                if converting:
                    pending.append(
                        (
                            record.row_number,
                            _cell(record, config.name_column),
                            _cell(record, config.phone_column),
                            reason,
                        )
                    )
                elif progress.invalid <= 5 and on_message is not None:
                    on_message(f"Lỗi dòng {record.row_number}: {reason}")
            else:
                for name, phone in contacts:
                    duplicated = bool(
                        config.dedupe and deduper is not None and deduper.is_duplicate(phone)
                    )
                    if duplicated:
                        progress.duplicate += 1
                        continue
                    if converting and exporter is not None:
                        finished = exporter.add(name, phone)
                        progress.exported += 1
                        progress.valid += 1
                        if finished is not None:
                            progress.files_created = exporter.files_created
                            progress.last_file = finished.name
                            logger.write(
                                f"Đã tạo {finished.name} ({config.contacts_per_file} liên hệ)."
                            )
                            file_closed = True
                    else:
                        progress.valid += 1
            progress.processed += 1
            rows_since_commit += 1
            next_offset = record.end_offset
            next_row = record.row_number + 1
            control.observe(progress)
            if file_closed or rows_since_commit >= config.chunk_size:
                durable()
                rows_since_commit = 0
            emit()
            if not control.keep_going():
                cancelled = True
                break

        if cancelled:
            if deduper is not None:
                deduper.rollback()
            pending.clear()
            if exporter is not None:
                exporter.flush()
                exporter.close()
            logger.write("Đã dừng. Có thể bấm Tiếp tục để chạy từ tiến trình đã lưu.")
            emit(force=True)
            return JobResult(True, "Đã dừng theo yêu cầu.", progress, None)

        if converting and exporter is not None and errors is not None:
            finished = exporter.finish()
            if finished is not None:
                progress.files_created = exporter.files_created
                progress.last_file = finished.name
                logger.write(f"Đã tạo {finished.name} ({_last_count(finished)} liên hệ).")
            if pending:
                _flush_errors(errors, pending)
                pending.clear()
            elif errors is not None:
                errors.flush()
            progress.files_created = exporter.files_created
            progress.elapsed_sec = elapsed()
            progress.total = progress.processed
            durable()
            report = Report(
                total_rows=progress.processed,
                valid=progress.valid,
                invalid=progress.invalid,
                duplicate=progress.duplicate,
                exported=progress.exported,
                vcf_files=progress.files_created,
                elapsed_sec=progress.elapsed_sec,
                average_speed=progress.speed,
            )
            (config.output_dir / "report.txt").write_text(report.render(), encoding="utf-8")
            guide = write_import_guide(config.output_dir)
            if guide is not None:
                logger.write(f"Đã ghi {guide.name}. Nhập các file VCF theo thứ tự trong file này.")
            logger.write(
                "Hoàn tất. "
                f"Hợp lệ {progress.valid}, lỗi {progress.invalid}, "
                f"trùng {progress.duplicate}, file {progress.files_created}."
            )
            emit(force=True)
            finished_cleanly = True
            return JobResult(False, "Chuyển đổi hoàn tất.", progress, report)

        if deduper is not None:
            deduper.commit()
        progress.elapsed_sec = elapsed()
        progress.total = progress.processed
        logger.write(
            "Kiểm tra xong. "
            f"Hợp lệ {progress.valid}, lỗi {progress.invalid}, trùng {progress.duplicate}."
        )
        emit(force=True)
        return JobResult(False, "Kiểm tra dữ liệu hoàn tất.", progress, None)
    except VcfStructureError:
        if deduper is not None:
            deduper.rollback()
        raise
    finally:
        logger.close()
        if errors is not None:
            errors.close()
        if exporter is not None:
            exporter.close()
        if deduper is not None:
            deduper.close()
        close = getattr(stream, "close", None) if "stream" in locals() else None
        if callable(close):
            close()
        if finished_cleanly:
            clear_resume_state(config.output_dir)
        if not converting:
            preview = config.output_dir / ".contact_to_vcf" / "preview-dedupe.sqlite"
            for extra in (preview, Path(str(preview) + "-wal"), Path(str(preview) + "-shm")):
                if extra.is_file():
                    extra.unlink()


def _estimate(config: JobConfig) -> int | None:
    if config.file_format == "xlsx":
        from openpyxl import load_workbook

        workbook = load_workbook(config.input_path, read_only=True, data_only=True)
        try:
            total = workbook.worksheets[0].max_row
        finally:
            workbook.close()
        if total is None:
            return None
        if config.has_header:
            return max(0, int(total) - 1)
        return int(total)
    total = estimate_text_rows(config.input_path)
    if config.has_header:
        return max(0, total - 1)
    return total


def _cell(record: RawRecord, index: int) -> str:
    if 0 <= index < len(record.columns):
        return record.columns[index]
    return ""


def _flush_errors(errors: ErrorLog, pending: list[tuple[int, str, str, str]]) -> int:
    errors.write_rows(pending)
    return errors.flush()


def _last_count(path: Path) -> int:
    _begin, end = count_markers(path)
    return end


def _make_checkpoint(
    config: JobConfig,
    size: int,
    fingerprint: str,
    progress: Progress,
    exporter: VcfExporter,
    error_bytes: int,
    next_offset: int,
    next_row: int,
) -> Checkpoint:
    return Checkpoint(
        version=CHECKPOINT_VERSION,
        input_path=str(config.input_path),
        input_size=size,
        input_fingerprint=fingerprint,
        file_format=config.file_format,
        delimiter=config.delimiter,
        has_header=config.has_header,
        name_column=config.name_column,
        phone_column=config.phone_column,
        contacts_per_file=config.contacts_per_file,
        dedupe=config.dedupe,
        normalize_phone=config.normalize_phone,
        vn_to_e164=config.vn_to_e164,
        keep_original_name=config.keep_original_name,
        split_directories=config.split_directories,
        files_per_batch=config.files_per_batch,
        encoding=config.encoding,
        chunk_size=config.chunk_size,
        next_offset=0 if config.file_format == "xlsx" else next_offset,
        next_row=next_row,
        processed=progress.processed,
        valid=progress.valid,
        invalid=progress.invalid,
        duplicate=progress.duplicate,
        exported=progress.exported,
        completed_files=exporter.completed_files,
        current_file_contacts=exporter.current_count,
        errors_bytes=error_bytes,
        elapsed_sec=progress.elapsed_sec,
        output_dir=str(config.output_dir),
        total_rows=progress.total,
        column_names=[],
    )


class _Logger:
    def __init__(self, path: Path, on_message) -> None:
        self._on_message = on_message
        path.parent.mkdir(parents=True, exist_ok=True)
        self._handle = path.open("a", encoding="utf-8")

    def write(self, message: str) -> None:
        stamp = time.strftime("%Y-%m-%d %H:%M:%S")
        self._handle.write(f"[{stamp}] {message}\n")
        self._handle.flush()
        if self._on_message is not None:
            self._on_message(message)

    def close(self) -> None:
        self._handle.close()
