"""Detect source format and build a streaming reader."""

from __future__ import annotations

import csv
from pathlib import Path

from models.records import FileInspection, JobConfig
from parsers.csv_parser import CsvReader
from parsers.txt_parser import TxtReader
from parsers.xlsx_parser import XlsxReader
from utils.fingerprint import estimate_text_rows, file_fingerprint

Reader = CsvReader | TxtReader | XlsxReader


def detect_format(path: Path) -> str:
    suffix = path.suffix.lower()
    if suffix == ".csv":
        return "csv"
    if suffix == ".xlsx":
        return "xlsx"
    if suffix == ".txt":
        return "txt"
    raise ValueError("Chỉ hỗ trợ file CSV, XLSX và TXT")


def detect_encoding(path: Path) -> str:
    sample = path.read_bytes()[:8192]
    try:
        sample.decode("utf-8-sig")
    except UnicodeDecodeError:
        return "cp1258"
    return "utf-8-sig"


def sniff_delimiter(path: Path, encoding: str) -> str:
    """Pick the delimiter that actually splits the sample into columns."""
    with path.open("rb") as handle:
        raw = handle.read(8192)
    try:
        sample = raw.decode(encoding)
    except UnicodeError:
        return ","
    counts = {delimiter: sample.count(delimiter) for delimiter in ",;\t|"}
    try:
        chosen = csv.Sniffer().sniff(sample, delimiters=",;\t|").delimiter
    except csv.Error:
        chosen = ""
    if chosen and counts.get(chosen, 0) > 0:
        return chosen
    best = max(counts, key=counts.get)
    if counts[best] > 0:
        return best
    return ","


def open_reader(config: JobConfig) -> Reader:
    if config.file_format == "csv":
        return CsvReader(
            config.input_path,
            config.delimiter,
            config.has_header,
            config.encoding,
        )
    if config.file_format == "txt":
        return TxtReader(
            config.input_path,
            config.delimiter,
            config.has_header,
            config.encoding,
        )
    if config.file_format == "xlsx":
        return XlsxReader(config.input_path, config.has_header)
    raise ValueError("Định dạng nguồn không được hỗ trợ")


def inspect_source(
    path: Path,
    *,
    file_format: str | None = None,
    delimiter: str | None = None,
    has_header: bool | None = None,
    encoding: str | None = None,
) -> FileInspection:
    resolved_format = file_format or detect_format(path)
    resolved_encoding = encoding or (
        "utf-8-sig" if resolved_format == "xlsx" else detect_encoding(path)
    )
    if resolved_format == "xlsx":
        resolved_delimiter = ""
        header = True if has_header is None else has_header
        reader: Reader = XlsxReader(path, header)
        estimated = _xlsx_estimate(path)
        if header and estimated:
            estimated = max(0, estimated - 1)
    else:
        header = (resolved_format != "txt") if has_header is None else has_header
        if delimiter in (None, "", "auto"):
            resolved_delimiter = "|" if resolved_format == "txt" else sniff_delimiter(path, resolved_encoding)
        else:
            resolved_delimiter = delimiter
        if resolved_format == "csv":
            reader = CsvReader(path, resolved_delimiter, header, resolved_encoding)
        else:
            reader = TxtReader(path, resolved_delimiter, header, resolved_encoding)
        estimated = estimate_text_rows(path)
        if header and estimated:
            estimated = max(0, estimated - 1)
    columns = reader.column_names()
    samples: list[tuple[str, ...]] = []
    for record in reader.iter_stream(0, 1):
        if record.error:
            continue
        samples.append(record.columns)
        if len(samples) >= 5:
            break
    size, fingerprint = file_fingerprint(path)
    return FileInspection(
        path=path,
        file_format=resolved_format,
        encoding=resolved_encoding,
        delimiter=resolved_delimiter,
        has_header=header,
        columns=columns,
        samples=samples,
        size=size,
        fingerprint=fingerprint,
        estimated_rows=estimated,
    )


def _xlsx_estimate(path: Path) -> int | None:
    from openpyxl import load_workbook

    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = workbook.worksheets[0]
        max_row = sheet.max_row
    finally:
        workbook.close()
    if max_row is None:
        return None
    return int(max_row)
