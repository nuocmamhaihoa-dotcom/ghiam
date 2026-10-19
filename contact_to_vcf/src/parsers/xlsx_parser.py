"""Streaming XLSX reader using openpyxl read-only mode."""

from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook

from models.records import RawRecord
from parsers.column_suggest import looks_like_phone


class XlsxReader:
    def __init__(self, path: Path, has_header: bool) -> None:
        self.path = path
        self.has_header = has_header
        self._first_loaded = False
        self._first: tuple[str, ...] | None = None

    def column_names(self) -> list[str]:
        first = self._first_values()
        if not first:
            return []
        cells = list(first)
        if self._skip_first(cells):
            return _header_names(cells)
        width = len(cells)
        while width > 0 and cells[width - 1].strip() == "":
            width -= 1
        if width == 0:
            return []
        return [f"Cột {index + 1}" for index in range(width)]

    def drops_first_row(self) -> bool:
        first = self._first_values()
        if not first:
            return False
        return self._skip_first(list(first))

    def iter_stream(self, start_offset: int, start_row_number: int):
        del start_offset
        workbook = load_workbook(self.path, read_only=True, data_only=True)
        try:
            next_row = 1
            for sheet in workbook.worksheets:
                first_row = True
                for row in sheet.iter_rows(values_only=True):
                    row_number = next_row
                    next_row += 1
                    columns = tuple(_cell_to_str(value) for value in row)
                    if first_row:
                        first_row = False
                        if self._skip_first(list(columns)):
                            continue
                    if row_number < start_row_number:
                        continue
                    if all(cell == "" for cell in columns):
                        continue
                    yield RawRecord(row_number, columns, row_number)
        finally:
            workbook.close()

    def _first_values(self) -> tuple[str, ...] | None:
        if self._first_loaded:
            return self._first
        self._first_loaded = True
        workbook = load_workbook(self.path, read_only=True, data_only=True)
        try:
            if not workbook.worksheets:
                self._first = None
                return None
            first = next(workbook.worksheets[0].iter_rows(values_only=True), None)
        finally:
            workbook.close()
        if first is None:
            self._first = None
            return None
        self._first = tuple(_cell_to_str(value) for value in first)
        return self._first

    def _skip_first(self, cells: list[str]) -> bool:
        """Drop a title row. Never drop a row that already contains a phone."""
        if any(looks_like_phone(cell) for cell in cells):
            return False
        return self.has_header


def _cell_to_str(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, int):
        return str(value)
    if isinstance(value, float):
        if value.is_integer():
            return str(int(value))
        return format(value, "f").rstrip("0").rstrip(".")
    return str(value)


def _header_names(cells: list[str]) -> list[str]:
    trimmed = list(cells)
    while trimmed and trimmed[-1].strip() == "":
        trimmed.pop()
    names: list[str] = []
    seen: dict[str, int] = {}
    for index, cell in enumerate(trimmed):
        base = cell.strip() or f"Cột {index + 1}"
        count = seen.get(base, 0) + 1
        seen[base] = count
        names.append(base if count == 1 else f"{base} ({count})")
    return names
