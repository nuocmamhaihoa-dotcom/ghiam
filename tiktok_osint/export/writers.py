from __future__ import annotations

import csv
import json
from datetime import datetime
from pathlib import Path
from typing import Any

from tiktok_osint.domain.dedupe import annotate_cross_profile
from tiktok_osint.errors import ExportError
from tiktok_osint.storage.db import init_db, make_engine, session_factory
from tiktok_osint.storage.orm import ExtractedContactRow, PublicProfileRow, ScanJobRow, ScanTargetRow
from tiktok_osint.storage.repo import Repository

FLAT_FIELDS = [
    "username",
    "nickname",
    "bio",
    "followers",
    "following",
    "likes",
    "verified",
    "private_account",
    "bio_link",
    "source_url",
    "target_status",
    "contact_kind",
    "contact_value",
    "contact_raw",
    "contact_source",
    "confidence",
    "evidence",
    "also_seen_on",
]

CONTACT_BOOK_FIELDS = ["id", "display_name", "phone_raw", "phone_e164", "email"]
SYNC_RESULT_FIELDS = [
    "tiktok_username",
    "display_name_shown",
    "contact_name",
    "contact_phone_e164",
    "contact_email",
    "matched_public_profile",
    "public_nickname",
    "public_followers",
    "public_verified",
    "public_profile_url",
    "note",
    "match_key",
]


def flat_rows(profiles: list[dict[str, Any]]) -> list[dict[str, str]]:
    rows: list[dict[str, str]] = []
    for profile in profiles:
        base = {
            "username": _s(profile.get("username")),
            "nickname": _s(profile.get("nickname")),
            "bio": _s(profile.get("bio")),
            "followers": _s(profile.get("followers")),
            "following": _s(profile.get("following")),
            "likes": _s(profile.get("likes")),
            "verified": _s(profile.get("verified")),
            "private_account": _s(profile.get("private_account")),
            "bio_link": _s(profile.get("bio_link")),
            "source_url": _s(profile.get("source_url")),
            "target_status": _s(profile.get("target_status")),
        }
        contacts = profile.get("contacts") or []
        if not contacts:
            rows.append(
                {
                    **base,
                    "contact_kind": "",
                    "contact_value": "",
                    "contact_raw": "",
                    "contact_source": "",
                    "confidence": "",
                    "evidence": "",
                }
            )
            continue
        for contact in contacts:
            rows.append(
                {
                    **base,
                    "contact_kind": _s(contact.get("kind")),
                    "contact_value": _s(contact.get("normalized_value")),
                    "contact_raw": _s(contact.get("raw_value")),
                    "contact_source": _s(contact.get("source")),
                    "confidence": _s(contact.get("confidence")),
                    "evidence": _s(contact.get("evidence")),
                }
            )
    return annotate_cross_profile(rows)


def write_csv(path: Path, rows: list[dict[str, str]]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=FLAT_FIELDS)
        writer.writeheader()
        writer.writerows(rows)


def write_xlsx(path: Path, profiles: list[dict[str, Any]], rows: list[dict[str, str]]) -> None:
    try:
        from openpyxl import Workbook
    except ImportError as exc:
        raise ExportError("Cài openpyxl để xuất Excel") from exc
    path.parent.mkdir(parents=True, exist_ok=True)
    book = Workbook()
    summary = book.active
    summary.title = "Profiles"
    profile_headers = [
        "username",
        "nickname",
        "bio",
        "followers",
        "following",
        "likes",
        "verified",
        "private_account",
        "bio_link",
        "target_status",
    ]
    summary.append(profile_headers)
    for profile in profiles:
        summary.append([profile.get(header) for header in profile_headers])
    contacts = book.create_sheet("Contacts")
    contacts.append(FLAT_FIELDS)
    for row in rows:
        contacts.append([row.get(field, "") for field in FLAT_FIELDS])
    book.save(path)


def write_contact_book_csv(path: Path, contacts: list[dict[str, Any]]) -> None:
    _write_dict_csv(path, CONTACT_BOOK_FIELDS, contacts)


def write_contact_book_xlsx(path: Path, name: str, contacts: list[dict[str, Any]]) -> None:
    book = _new_workbook()
    sheet = book.active
    sheet.title = "Contacts"
    sheet.append(["Danh bạ", *CONTACT_BOOK_FIELDS])
    for contact in contacts:
        sheet.append([name, *[contact.get(field, "") for field in CONTACT_BOOK_FIELDS]])
    book.save(path)


def write_sync_results_csv(path: Path, rows: list[dict[str, Any]]) -> None:
    _write_dict_csv(path, SYNC_RESULT_FIELDS, [_flatten_sync_result(row) for row in rows])


def write_sync_results_xlsx(path: Path, session: dict[str, Any], rows: list[dict[str, Any]]) -> None:
    book = _new_workbook()
    summary = book.active
    summary.title = "Session"
    summary.append(["id", "book_id", "status", "note", "recorded_count", "created_at", "updated_at"])
    summary.append(
        [
            session.get("id"),
            session.get("book_id"),
            session.get("status"),
            session.get("note"),
            session.get("recorded_count"),
            session.get("created_at"),
            session.get("updated_at"),
        ]
    )
    results = book.create_sheet("Displayed accounts")
    results.append(SYNC_RESULT_FIELDS)
    for row in rows:
        flat = _flatten_sync_result(row)
        results.append([flat.get(field, "") for field in SYNC_RESULT_FIELDS])
    book.save(path)


def write_sqlite(path: Path, repo: Repository, job_id: str) -> None:
    job = repo.get_job(job_id)
    if job is None:
        from tiktok_osint.errors import JobNotFound

        raise JobNotFound(job_id)
    targets = repo.list_targets(job_id)
    profiles = [profile for profile in repo.profiles_for_job(job_id) if profile.get("id")]
    if path.exists():
        path.unlink()
    engine = make_engine(f"sqlite:///{path}")
    init_db(engine)
    sessions = session_factory(engine)
    with sessions() as session:
        session.add(
            ScanJobRow(
                id=str(job["id"]),
                name=str(job["name"]),
                status=str(job["status"]),
                checkpoint_json=json.dumps(job["checkpoint"], ensure_ascii=False),
                error=job["error"] if isinstance(job["error"], str) else None,
                created_at=_parse_dt(str(job["created_at"])),
                updated_at=_parse_dt(str(job["updated_at"])),
            )
        )
        for target in targets:
            session.add(
                ScanTargetRow(
                    id=str(target["id"]),
                    job_id=str(target["job_id"]),
                    raw_input=str(target["raw_input"]),
                    username=target["username"],
                    profile_url=target["profile_url"],
                    status=str(target["status"]),
                    attempts=int(target["attempts"]),
                    last_error=target["last_error"],
                    created_at=_parse_dt(str(target["created_at"])),
                    updated_at=_parse_dt(str(target["updated_at"])),
                )
            )
        for profile in profiles:
            session.add(
                PublicProfileRow(
                    id=str(profile["id"]),
                    username=str(profile["username"]),
                    nickname=profile.get("nickname"),
                    bio=profile.get("bio"),
                    followers=profile.get("followers"),
                    following=profile.get("following"),
                    likes=profile.get("likes"),
                    verified=bool(profile.get("verified")),
                    private_account=bool(profile.get("private_account")),
                    avatar_url=profile.get("avatar_url"),
                    bio_link=profile.get("bio_link"),
                    source_url=str(profile.get("source_url") or ""),
                    partial=bool(profile.get("partial")),
                    job_id=profile.get("job_id"),
                    scraped_at=_parse_dt(str(profile["scraped_at"])),
                )
            )
            for contact in profile.get("contacts") or []:
                session.add(
                    ExtractedContactRow(
                        id=str(contact["id"]),
                        profile_id=str(contact["profile_id"]),
                        kind=str(contact["kind"]),
                        raw_value=str(contact["raw_value"]),
                        normalized_value=str(contact["normalized_value"]),
                        source=str(contact["source"]),
                        confidence=float(contact["confidence"]),
                        evidence=contact.get("evidence"),
                    )
                )
        session.commit()


def _s(value: object) -> str:
    if value is None:
        return ""
    return str(value)


def _write_dict_csv(path: Path, fields: list[str], rows: list[dict[str, Any]]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=fields, extrasaction="ignore")
        writer.writeheader()
        writer.writerows(rows)


def _new_workbook() -> Any:
    try:
        from openpyxl import Workbook
    except ImportError as exc:
        raise ExportError("Cài openpyxl để xuất Excel") from exc
    return Workbook()


def _flatten_sync_result(row: dict[str, Any]) -> dict[str, Any]:
    contact = row.get("contact") if isinstance(row.get("contact"), dict) else {}
    profile = row.get("public_profile") if isinstance(row.get("public_profile"), dict) else {}
    username = str(row.get("tiktok_username") or "")
    return {
        "tiktok_username": username,
        "display_name_shown": row.get("display_name_shown"),
        "contact_name": contact.get("display_name"),
        "contact_phone_e164": contact.get("phone_e164"),
        "contact_email": contact.get("email"),
        "matched_public_profile": row.get("matched_public_profile"),
        "public_nickname": profile.get("nickname"),
        "public_followers": profile.get("followers"),
        "public_verified": profile.get("verified"),
        "public_profile_url": f"https://www.tiktok.com/@{username}" if username else "",
        "note": row.get("note"),
        "match_key": row.get("match_key"),
    }


def _parse_dt(value: str) -> datetime:
    return datetime.fromisoformat(value)
