"""Nhận cột từ nhiều form Excel rồi chuẩn hoá vào kho tổng.

Cách đã chọn: hybrid xác định — từ điển tiêu đề + hình dạng ô + nhớ form.
Không dùng LLM: cùng một file luôn ra cùng một map, chạy được khi mất mạng,
và chịu được file lớn.
"""

from __future__ import annotations

import csv
import hashlib
import io
import re
import unicodedata
from dataclasses import dataclass
from pathlib import Path
from typing import Any, Iterable

FIELDS = ("name", "phone", "address", "uid", "id")
FIELD_LABELS = {
    "name": "Tên",
    "phone": "Số điện thoại",
    "address": "Địa chỉ",
    "uid": "UID",
    "id": "ID",
}
MAX_SCAN_ROWS = 40
MAX_HEADER_SCAN = 25
SAMPLE_LIMIT = 20
HEADER_WEIGHT = 0.62
ASSIGN_MIN = 0.22

_PUNCT = re.compile(r"[^\w\s]+", re.UNICODE)
_SPACES = re.compile(r"\s+")
_PHONE_KEEP = re.compile(r"[^\d+]")
_UID_OK = re.compile(r"^\d{8,20}$")
_ID_OK = re.compile(r"^[A-Za-z0-9][A-Za-z0-9._/-]{0,39}$")
_HAS_LETTER = re.compile(r"[A-Za-zÀ-ỹ]")
_ADDRESS_MARK = re.compile(
    r"\b(duong|pho|quan|huyen|tinh|phuong|xa|thon|ap|ngo|hem|tp|tphcm|hn|"
    r"thanh pho|so nha|dia chi|q\d+|p\d+)\b"
    r"|[/]|(\d+\s*(duong|pho|ngo|hem))",
    re.IGNORECASE,
)
_STREET_WORDS = (
    "đường",
    "phố",
    "quận",
    "huyện",
    "tỉnh",
    "phường",
    "xã",
    "thôn",
    "ấp",
    "ngõ",
    "hẻm",
    "tp.",
    "tphcm",
    "thành phố",
    "số nhà",
    "địa chỉ",
)


def fold(text: str) -> str:
    raw = unicodedata.normalize("NFD", str(text or ""))
    stripped = "".join(ch for ch in raw if unicodedata.category(ch) != "Mn")
    folded = _PUNCT.sub(" ", stripped.casefold())
    return _SPACES.sub(" ", folded).strip()


def cell_text(value: object) -> str:
    """Đọc ô Excel thành chữ, tránh mất số 0 và tránh khoa học với số nguyên."""
    if value is None:
        return ""
    if isinstance(value, bool):
        return ""
    if isinstance(value, int):
        return str(value)
    if isinstance(value, float):
        if value != value:  # NaN
            return ""
        if abs(value) >= 2**53:
            return format(value, ".0f")
        if value.is_integer():
            return str(int(value))
        return format(value, ".15g")
    text = str(value).replace("\ufeff", "").strip()
    if text.lower() in {"none", "null", "nan"}:
        return ""
    return text


_HEADER_ALIASES: dict[str, frozenset[str]] = {
    "name": frozenset(
        {
            "name",
            "fullname",
            "full name",
            "customer name",
            "ten",
            "ho ten",
            "ho va ten",
            "hoten",
            "ten khach",
            "ten khach hang",
            "ten nguoi",
            "ten lien he",
            "ho ten khach",
            "customer",
            "khach hang",
        }
    ),
    "phone": frozenset(
        {
            "phone",
            "phone number",
            "tel",
            "telephone",
            "mobile",
            "cellphone",
            "sdt",
            "so dt",
            "so dien thoai",
            "dien thoai",
            "dt",
            "sdt khach",
            "so may",
            "so lien he",
        }
    ),
    "address": frozenset(
        {
            "address",
            "addr",
            "dia chi",
            "diachi",
            "dc",
            "dia chi nha",
            "noi o",
            "dia chi khach",
            "dia chi giao",
            "home address",
        }
    ),
    "uid": frozenset(
        {
            "uid",
            "facebook uid",
            "fb uid",
            "fbid",
            "fb id",
            "facebook id",
            "uid fb",
            "user uid",
            "psid",
        }
    ),
    "id": frozenset(
        {
            "id",
            "ma",
            "ma khach",
            "ma khach hang",
            "customer id",
            "ma so",
            "code",
            "id khach",
            "external id",
            "ref",
            "ma gd",
        }
    ),
}

_NAME_LAST = frozenset({"ho", "ho dem", "last name", "surname", "family name"})
_NAME_FIRST = frozenset({"ten rieng", "first name", "given name"})
_SKIP_HEADERS = frozenset(
    {
        "stt",
        "so thu tu",
        "no",
        "tt",
        "ghi chu",
        "note",
        "notes",
        "email",
        "mail",
        "ngay",
        "date",
        "ngay sinh",
    }
)


def header_field(text: str) -> str | None:
    key = fold(text)
    if not key or key in _SKIP_HEADERS:
        return None
    if key in _NAME_LAST or key in _NAME_FIRST:
        return "name"
    for field_name, aliases in _HEADER_ALIASES.items():
        if key in aliases:
            return field_name
    for field_name, aliases in _HEADER_ALIASES.items():
        if any(alias and (key == alias or key.startswith(alias + " ") or alias in key.split()) for alias in aliases):
            if len(key) <= 28:
                return field_name
    return None


def header_score(text: str) -> dict[str, float]:
    scores = {name: 0.0 for name in FIELDS}
    key = fold(text)
    if not key or key in _SKIP_HEADERS:
        return scores
    hit = header_field(text)
    if hit:
        scores[hit] = 1.0 if key in _HEADER_ALIASES[hit] or key in _NAME_LAST | _NAME_FIRST else 0.82
        return scores
    for field_name, aliases in _HEADER_ALIASES.items():
        for alias in aliases:
            if alias and alias in key:
                scores[field_name] = max(scores[field_name], 0.55 if len(alias) >= 3 else 0.0)
    return scores


def _vn_mobile(digits: str) -> str:
    if digits.startswith("84") and len(digits) == 11:
        return "0" + digits[2:]
    if len(digits) == 9 and digits[0] in "35789":
        return "0" + digits
    if len(digits) == 10 and digits.startswith(("02", "03", "05", "07", "08", "09")):
        return digits
    if len(digits) == 11 and digits.startswith("02"):
        return digits
    return ""


def normalize_phone(raw: str) -> str:
    """Chỉ nhận số Việt / +quốc tế. Số 12–20 chữ số để dành cho UID."""
    text = cell_text(raw)
    if text.lower().startswith("tel:"):
        text = text[4:].strip()
    compact = _PHONE_KEEP.sub("", text)
    plus = compact.startswith("+")
    digits = "".join(ch for ch in compact if ch.isdigit())
    local = _vn_mobile(digits)
    if local:
        return local
    if plus and 8 <= len(digits) <= 15:
        return "+" + digits
    return ""


def phone_key(phone: str) -> str:
    return "".join(ch for ch in phone if ch.isdigit())


def normalize_uid(raw: str) -> str:
    text = cell_text(raw).replace(" ", "")
    if text.endswith(".0") and text[:-2].isdigit():
        text = text[:-2]
    if _UID_OK.match(text):
        if normalize_phone(text) and text.startswith(("03", "05", "07", "08", "09", "84")):
            return ""
        return text
    return ""


def normalize_id(raw: str) -> str:
    text = cell_text(raw)
    if not text:
        return ""
    if normalize_phone(text):
        return ""
    folded = text.replace(" ", "")
    if _ID_OK.match(folded) and not folded.isdigit():
        return folded[:40]
    if folded.isdigit() and 3 <= len(folded) <= 20 and not normalize_uid(folded):
        return folded
    return folded[:40] if _HAS_LETTER.search(folded) and len(folded) <= 40 else ""


def normalize_name(raw: str) -> str:
    text = _SPACES.sub(" ", cell_text(raw)).strip()
    if not text or normalize_phone(text):
        return ""
    if len(text) > 80:
        text = text[:80].rstrip()
    return text


def normalize_address(raw: str) -> str:
    text = _SPACES.sub(" ", cell_text(raw)).strip()
    if not text or normalize_phone(text):
        return ""
    return text[:200]


def _looks_name(text: str) -> bool:
    value = normalize_name(text)
    if not value:
        return False
    if any(mark in value.casefold() for mark in _STREET_WORDS):
        return False
    words = value.split()
    if not (1 <= len(words) <= 6):
        return False
    digits = sum(ch.isdigit() for ch in value)
    return digits <= 1 and bool(_HAS_LETTER.search(value))


def _looks_address(text: str) -> bool:
    value = cell_text(text)
    if len(value) < 10:
        return False
    folded = fold(value)
    if _ADDRESS_MARK.search(folded):
        return True
    return len(value) >= 18 and bool(_HAS_LETTER.search(value)) and not normalize_phone(value)


def is_index_column(values: list[str]) -> bool:
    numbers: list[int] = []
    for raw in values:
        text = cell_text(raw)
        if text.isdigit():
            numbers.append(int(text))
        elif text:
            return False
    if len(numbers) < 4:
        return False
    step = sum(1 for a, b in zip(numbers, numbers[1:]) if b - a == 1)
    return step / max(1, len(numbers) - 1) >= 0.7 and numbers[0] <= 5


def shape_scores(values: list[str]) -> dict[str, float]:
    scores = {name: 0.0 for name in FIELDS}
    useful = [cell_text(item) for item in values if cell_text(item)]
    if not useful:
        return scores
    if is_index_column(useful):
        return scores
    n = len(useful)
    phone_hits = sum(1 for item in useful if normalize_phone(item))
    uid_hits = sum(1 for item in useful if normalize_uid(item))
    id_hits = sum(1 for item in useful if normalize_id(item))
    name_hits = sum(1 for item in useful if _looks_name(item))
    address_hits = sum(1 for item in useful if _looks_address(item))
    scores["phone"] = phone_hits / n
    scores["uid"] = uid_hits / n
    scores["id"] = id_hits / n
    scores["name"] = name_hits / n
    scores["address"] = address_hits / n
    if scores["phone"] >= 0.7:
        scores["uid"] *= 0.15
        scores["id"] *= 0.2
    if scores["uid"] >= 0.7:
        scores["phone"] *= 0.15
        scores["id"] *= 0.25
    return scores


@dataclass
class ColumnGuess:
    index: int
    header: str
    letter: str
    field: str | None
    confidence: float
    header_scores: dict[str, float]
    shape_scores: dict[str, float]


@dataclass
class DetectedSheet:
    name: str
    header_row: int
    mapping: dict[str, int]
    columns: list[ColumnGuess]
    sample: list[dict[str, str]]
    row_count: int
    confidence: float
    warnings: list[str]
    method: str
    fingerprint: str
    reused_template: bool = False


@dataclass
class PreviewResult:
    sheets: list[DetectedSheet]
    chosen: str
    filename: str


def column_letter(index: int) -> str:
    if index < 0:
        return "?"
    n = index + 1
    letters = []
    while n:
        n, rem = divmod(n - 1, 26)
        letters.append(chr(65 + rem))
    return "".join(reversed(letters))


def combine_scores(header: dict[str, float], shape: dict[str, float]) -> dict[str, float]:
    header_peak = max(header.values()) if header else 0.0
    weight_h = HEADER_WEIGHT if header_peak >= 0.55 else 0.2
    weight_s = 1.0 - weight_h
    return {
        name: round(header.get(name, 0.0) * weight_h + shape.get(name, 0.0) * weight_s, 4)
        for name in FIELDS
    }


def assign_fields(columns: list[ColumnGuess]) -> dict[str, int]:
    mapping: dict[str, int] = {}
    taken: set[int] = set()
    candidates: list[tuple[float, str, int]] = []
    for col in columns:
        combined = combine_scores(col.header_scores, col.shape_scores)
        for field_name, score in combined.items():
            if score >= ASSIGN_MIN:
                candidates.append((score, field_name, col.index))
    candidates.sort(key=lambda item: (-item[0], item[1], item[2]))
    for score, field_name, index in candidates:
        if field_name in mapping or index in taken:
            continue
        mapping[field_name] = index
        taken.add(index)
        for col in columns:
            if col.index == index:
                col.field = field_name
                col.confidence = score
    return mapping


def detect_header_row(rows: list[list[object]]) -> tuple[int, str]:
    best_row = 0
    best_score = -1.0
    method = "shape"
    limit = min(len(rows), MAX_HEADER_SCAN)
    for index in range(limit):
        cells = [cell_text(item) for item in rows[index]]
        if not any(cells):
            continue
        hits = sum(1 for cell in cells if header_field(cell))
        if hits < 2:
            continue
        following = rows[index + 1 : index + 6]
        data_like = 0
        for follow in following:
            texts = [cell_text(item) for item in follow]
            if any(normalize_phone(item) or _looks_name(item) or normalize_uid(item) for item in texts):
                data_like += 1
        score = hits + data_like * 0.4
        if score > best_score:
            best_score = score
            best_row = index
            method = "header"
    if best_score < 0:
        return 0, "shape"
    return best_row, method


def fingerprint_for(headers: Iterable[str]) -> str:
    joined = "|".join(fold(item) for item in headers)
    return hashlib.sha256(joined.encode("utf-8")).hexdigest()[:24]


def rows_from_csv(payload: bytes) -> list[list[object]]:
    for encoding in ("utf-8-sig", "utf-8", "cp1258", "cp1252", "latin-1"):
        try:
            text = payload.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    else:
        text = payload.decode("utf-8", errors="replace")
    sample = text[:4000]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel
    reader = csv.reader(io.StringIO(text), dialect)
    return [list(row) for row in reader]


def rows_from_xlsx(payload: bytes, sheet: str | None = None) -> tuple[str, list[str], list[list[object]]]:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise RuntimeError("Thiếu openpyxl để đọc file Excel") from exc
    book = load_workbook(io.BytesIO(payload), read_only=True, data_only=True)
    try:
        names = list(book.sheetnames)
        chosen = sheet if sheet in names else names[0]
        table: list[list[object]] = []
        for row in book[chosen].iter_rows(values_only=True):
            table.append(list(row))
        return chosen, names, table
    finally:
        book.close()


def list_xlsx_sheets(payload: bytes) -> list[str]:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise RuntimeError("Thiếu openpyxl để đọc file Excel") from exc
    book = load_workbook(io.BytesIO(payload), read_only=True, data_only=True)
    try:
        return list(book.sheetnames)
    finally:
        book.close()


def detect_table(
    rows: list[list[object]],
    sheet_name: str,
    saved_map: dict[str, int] | None = None,
) -> DetectedSheet:
    if not rows:
        return DetectedSheet(
            name=sheet_name,
            header_row=0,
            mapping={},
            columns=[],
            sample=[],
            row_count=0,
            confidence=0.0,
            warnings=["Trang tính trống"],
            method="empty",
            fingerprint="",
        )
    width = max((len(row) for row in rows), default=0)
    padded = [list(row) + [""] * (width - len(row)) for row in rows]
    header_row, method = detect_header_row(padded)
    headers = [cell_text(item) for item in padded[header_row]] if method == "header" else [""] * width
    data_start = header_row + 1 if method == "header" else 0
    data_rows = padded[data_start:]
    nonempty = [row for row in data_rows if any(cell_text(item) for item in row)]
    columns: list[ColumnGuess] = []
    for index in range(width):
        sample_values = [cell_text(row[index]) for row in nonempty[:MAX_SCAN_ROWS]]
        guess = ColumnGuess(
            index=index,
            header=headers[index] if index < len(headers) else "",
            letter=column_letter(index),
            field=None,
            confidence=0.0,
            header_scores=header_score(headers[index] if index < len(headers) else ""),
            shape_scores=shape_scores(sample_values),
        )
        columns.append(guess)
    reused = False
    if saved_map:
        mapping = {field_name: index for field_name, index in saved_map.items() if 0 <= index < width}
        for col in columns:
            for field_name, index in mapping.items():
                if col.index == index:
                    col.field = field_name
                    col.confidence = 1.0
        reused = True
        method = "template"
    else:
        mapping = assign_fields(columns)
    records = apply_mapping(nonempty, mapping)
    warnings: list[str] = []
    if "name" not in mapping and "phone" not in mapping:
        warnings.append("Chưa chắc cột tên và số điện thoại. Hãy sửa map trước khi nạp.")
    if any(col.shape_scores.get("uid", 0) > 0 and col.field == "uid" for col in columns):
        if any(len(cell_text(row[mapping["uid"]])) >= 16 for row in nonempty[:MAX_SCAN_ROWS] if "uid" in mapping):
            warnings.append("UID dài: nếu Excel lưu dạng số có thể mất chính xác. Nên để cột UID là chữ.")
    mapped_conf = [col.confidence for col in columns if col.field]
    confidence = round(sum(mapped_conf) / len(mapped_conf), 3) if mapped_conf else 0.0
    if reused:
        confidence = max(confidence, 0.95)
    return DetectedSheet(
        name=sheet_name,
        header_row=header_row,
        mapping=mapping,
        columns=columns,
        sample=records[:SAMPLE_LIMIT],
        row_count=len(records),
        confidence=confidence,
        warnings=warnings,
        method=method,
        fingerprint=fingerprint_for(headers),
        reused_template=reused,
    )


def apply_mapping(rows: list[list[object]], mapping: dict[str, int]) -> list[dict[str, str]]:
    records: list[dict[str, str]] = []
    for row in rows:
        record = {name: "" for name in FIELDS}
        if "name" in mapping and mapping["name"] < len(row):
            record["name"] = normalize_name(cell_text(row[mapping["name"]]))
        if "phone" in mapping and mapping["phone"] < len(row):
            record["phone"] = normalize_phone(cell_text(row[mapping["phone"]]))
        if "address" in mapping and mapping["address"] < len(row):
            record["address"] = normalize_address(cell_text(row[mapping["address"]]))
        if "uid" in mapping and mapping["uid"] < len(row):
            record["uid"] = normalize_uid(cell_text(row[mapping["uid"]]))
        if "id" in mapping and mapping["id"] < len(row):
            record["id"] = normalize_id(cell_text(row[mapping["id"]]))
        if not any(record.values()):
            continue
        if not record["name"] and not record["phone"] and not record["uid"] and not record["id"]:
            continue
        records.append(record)
    return records


def preview_bytes(
    payload: bytes,
    filename: str,
    templates: dict[str, dict[str, int]] | None = None,
    sheet: str | None = None,
) -> PreviewResult:
    templates = templates or {}
    suffix = Path(filename).suffix.lower()
    sheets: list[DetectedSheet] = []
    if suffix in {".xlsx", ".xlsm"}:
        names = list_xlsx_sheets(payload)
        targets = [sheet] if sheet and sheet in names else names
        for name in targets:
            _, _, rows = rows_from_xlsx(payload, name)
            headers = []
            if rows:
                header_row, method = detect_header_row(rows)
                if method == "header":
                    headers = [cell_text(item) for item in rows[header_row]]
            saved = templates.get(fingerprint_for(headers))
            sheets.append(detect_table(rows, name, saved))
    elif suffix in {".csv", ".tsv", ".txt"}:
        rows = rows_from_csv(payload)
        header_row, method = detect_header_row(rows)
        headers = [cell_text(item) for item in rows[header_row]] if rows and method == "header" else []
        saved = templates.get(fingerprint_for(headers))
        sheets.append(detect_table(rows, "csv", saved))
    else:
        raise ValueError("Chỉ nhận .xlsx, .xlsm, .csv, .tsv, .txt")
    if not sheets:
        raise ValueError("Không đọc được trang tính")
    chosen = max(sheets, key=lambda item: (item.confidence, item.row_count))
    return PreviewResult(sheets=sheets, chosen=chosen.name, filename=filename)


def sheet_to_dict(sheet: DetectedSheet) -> dict[str, Any]:
    return {
        "name": sheet.name,
        "headerRow": sheet.header_row,
        "mapping": sheet.mapping,
        "columns": [
            {
                "index": col.index,
                "letter": col.letter,
                "header": col.header,
                "field": col.field,
                "confidence": col.confidence,
                "headerScores": col.header_scores,
                "shapeScores": col.shape_scores,
            }
            for col in sheet.columns
        ],
        "sample": sheet.sample,
        "rowCount": sheet.row_count,
        "confidence": sheet.confidence,
        "warnings": sheet.warnings,
        "method": sheet.method,
        "fingerprint": sheet.fingerprint,
        "reusedTemplate": sheet.reused_template,
        "fields": FIELD_LABELS,
    }


def records_from_bytes(
    payload: bytes,
    filename: str,
    sheet: str | None,
    mapping: dict[str, int],
    header_row: int | None = None,
    has_header: bool | None = None,
) -> list[dict[str, str]]:
    suffix = Path(filename).suffix.lower()
    if suffix in {".xlsx", ".xlsm"}:
        _, _, rows = rows_from_xlsx(payload, sheet)
    else:
        rows = rows_from_csv(payload)
    if not rows:
        return []
    if header_row is None or has_header is None:
        header_row, method = detect_header_row(rows)
        has_header = method == "header"
    start = header_row + 1 if has_header else 0
    width = max((len(row) for row in rows), default=0)
    padded = [list(row) + [""] * (width - len(row)) for row in rows[start:]]
    nonempty = [row for row in padded if any(cell_text(item) for item in row)]
    return apply_mapping(nonempty, mapping)
